from datetime import date, datetime
from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.gis.geos import Point
from django.db import transaction
from django.db.models import Count, OuterRef, Q, Subquery
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from assets.models import AssetType, MunicipalAsset
from complaints.models import Complaint
from locations.models import Neighborhood
from maintenance.models import WorkOrder
from municipal_services.models import MunicipalService
from notifications.models import Notification

from .decorators import municipal_required
from .forms import AssetExcelUploadForm, MunicipalAssetForm
from .asset_etl import AssetETLError, process_asset_excel

User = get_user_model()

@municipal_required
def dashboard(request):
    complaints = Complaint.objects.all()

    recent_complaints = (
        complaints
        .select_related(
            "citizen",
            "service",
            "category",
            "neighborhood",
        )
        .order_by("-created_at")[:8]
    )

    context = {
        "total_count": complaints.count(),

        "submitted_count": complaints.filter(
            status=Complaint.Status.SUBMITTED,
        ).count(),

        "under_review_count": complaints.filter(
            status=Complaint.Status.UNDER_REVIEW,
        ).count(),

        "assigned_count": complaints.filter(
            status=Complaint.Status.ASSIGNED,
        ).count(),

        "in_progress_count": complaints.filter(
            status=Complaint.Status.IN_PROGRESS,
        ).count(),

        "resolved_count": complaints.filter(
            status=Complaint.Status.RESOLVED,
        ).count(),

        "closed_count": complaints.filter(
            status=Complaint.Status.CLOSED,
        ).count(),

        "active_work_orders_count": (
            WorkOrder.objects
            .exclude(
                status__in=[
                    WorkOrder.Status.COMPLETED,
                    WorkOrder.Status.CANCELLED,
                ]
            )
            .count()
        ),

        "recent_complaints": recent_complaints,
    }

    return render(
        request,
        "municipal_portal/dashboard.html",
        context,
    )


@municipal_required
def services_assets(request):
    services = (
        MunicipalService.objects
        .prefetch_related("categories")
        .annotate(
            category_count=Count(
                "categories",
                distinct=True,
            ),
            asset_type_count=Count(
                "asset_types",
                distinct=True,
            ),
            asset_count=Count(
                "asset_types__assets",
                distinct=True,
            ),
        )
        .order_by("name_ar")
    )

    last_completed_maintenance = (
        WorkOrder.objects
        .filter(
            complaint__asset=OuterRef("pk"),
            status=WorkOrder.Status.COMPLETED,
            completed_at__isnull=False,
        )
        .order_by("-completed_at")
        .values("completed_at")[:1]
    )

    assets = (
        MunicipalAsset.objects
        .select_related(
            "asset_type",
            "asset_type__service",
            "neighborhood",
        )
        .annotate(
            last_maintenance_at=Subquery(
                last_completed_maintenance,
            )
        )
        .order_by(
            "asset_type__service__name_ar",
            "asset_code",
        )
    )
    assets_by_service = {}

    for asset in assets:
        service_id = asset.asset_type.service_id

        assets_by_service.setdefault(
            service_id,
            [],
        ).append(asset)

    service_sections = []

    for service in services:
        service_sections.append(
            {
                "service": service,
                "assets": assets_by_service.get(
                    service.pk,
                    [],
                ),
            }
        )
        editing_asset = None
    asset_form = MunicipalAssetForm()

    edit_asset_id = request.GET.get(
        "edit",
        "",
    ).strip()

    if edit_asset_id.isdigit():
        editing_asset = get_object_or_404(
            MunicipalAsset.objects.select_related(
                "asset_type",
                "asset_type__service",
            ),
            pk=int(edit_asset_id),
        )

        asset_form = MunicipalAssetForm(
            instance=editing_asset,
        )    

    return render(
        request,
        "municipal_portal/services_assets.html",
        {
            "services": services,
            "assets": assets,
            "service_sections": service_sections,
            "asset_form": asset_form,
"editing_asset": editing_asset,
            "excel_form": AssetExcelUploadForm(),
        },
    )

@municipal_required
@require_POST
def asset_create(request):
    form = MunicipalAssetForm(
        request.POST,
    )

    if form.is_valid():
        asset = form.save()

        messages.success(
            request,
            (
                f"تمت إضافة الأصل "
                f"{asset.asset_code} بنجاح."
            ),
        )

        return redirect(
            "municipal_portal:services_assets",
        )

    error_messages = []

    for field_name, field_errors in form.errors.items():
        if field_name == "__all__":
            field_label = ""
        else:
            field = form.fields.get(field_name)
            field_label = (
                field.label
                if field is not None
                else field_name
        )

    for error in field_errors:
        if field_label:
            error_messages.append(
                f"{field_label}: {error}"
            )
        else:
            error_messages.append(
                str(error)
            )

    messages.error(
    request,
    (
        "تعذر حفظ الأصل. يرجى تصحيح البيانات التالية: "
        + " • ".join(error_messages)
    ),
)

    return redirect(
        "municipal_portal:services_assets",
    )

@municipal_required
@require_POST
def asset_update(request, asset_id):
    asset = get_object_or_404(
        MunicipalAsset,
        pk=asset_id,
    )

    form = MunicipalAssetForm(
        request.POST,
        instance=asset,
    )

    if form.is_valid():
        updated_asset = form.save()

        messages.success(
            request,
            (
                f"تم تحديث الأصل "
                f"{updated_asset.asset_code} بنجاح."
            ),
        )

        return redirect(
            "municipal_portal:services_assets"
        )

    messages.error(
        request,
        "تعذر تحديث الأصل. تحققي من الحقول المدخلة.",
    )

    return redirect(
        f"{reverse('municipal_portal:services_assets')}"
        f"?edit={asset.pk}"
    )
@municipal_required
@require_POST
def asset_create(request):
    form = MunicipalAssetForm(
        request.POST,
    )

    if form.is_valid():
        asset = form.save()

        messages.success(
            request,
            (
                f"تمت إضافة الأصل "
                f"{asset.asset_code} بنجاح."
            ),
        )

        return redirect(
            "municipal_portal:services_assets",
        )

    error_messages = []

    for field_errors in form.errors.values():
        for error in field_errors:
            error_messages.append(str(error))

    messages.error(
        request,
        "تعذر إضافة الأصل: "
        + "، ".join(error_messages),
    )

    return redirect(
        "municipal_portal:services_assets",
    )


@municipal_required
@require_POST
def asset_import(request):
    form = AssetExcelUploadForm(
        request.POST,
        request.FILES,
    )

    if not form.is_valid():
        error_messages = []

        for field_errors in form.errors.values():
            for error in field_errors:
                error_messages.append(str(error))

        messages.error(
            request,
            (
                "تعذر رفع ملف Excel. "
                + (
                    " ".join(error_messages)
                    if error_messages
                    else "يرجى اختيار ملف صالح بصيغة xlsx."
                )
            ),
        )

        return redirect(
            "municipal_portal:services_assets",
        )

    try:
        result = process_asset_excel(
            form.cleaned_data["excel_file"]
        )

    except AssetETLError as error:
        messages.error(
            request,
            str(error),
        )

        return redirect(
            "municipal_portal:services_assets",
        )

    except Exception as error:
        messages.error(
            request,
        (
                "حدث خطأ أثناء معالجة ملف Excel: "
                f"{type(error).__name__}: {error}"
        ),
    )

        return redirect(
            "municipal_portal:services_assets",
    )

        return redirect(
            "municipal_portal:services_assets",
        )

    if result.processed == 0:
        messages.warning(
            request,
            (
                "لم يعثر النظام على صفوف بيانات داخل الملف. "
                f"تم تجاهل {result.empty_rows} صف فارغ."
            ),
        )

        return redirect(
            "municipal_portal:services_assets",
        )

    if result.added > 0:
        messages.success(
            request,
            f"تمت إضافة {result.added} أصل جديد بنجاح.",
        )

    if result.updated > 0:
        messages.success(
            request,
            f"تم تحديث بيانات {result.updated} أصل بنجاح.",
        )

    if result.unchanged > 0:
        messages.info(
            request,
            (
                f"لم تتغير بيانات {result.unchanged} أصل، "
                "لذلك لم يحتج النظام إلى تحديثها."
            ),
        )

    if result.empty_rows > 0:
        messages.info(
            request,
            (
                f"تم تجاهل {result.empty_rows} "
                "صف فارغ في ملف Excel."
            ),
        )

    if result.rejected > 0:
        messages.warning(
            request,
            (
                f"تعذر استيراد {result.rejected} صف. "
                "تم استيراد بقية الصفوف الصحيحة."
            ),
        )

        rejected_rows = [
            row
            for row in result.rows
            if row.status == "rejected"
        ]

        for row in rejected_rows[:10]:
            asset_description = (
                f" للأصل «{row.asset_code}»"
                if row.asset_code
                else ""
            )

            messages.error(
                request,
                (
                    f"الصف {row.row_number}"
                    f"{asset_description}: "
                    f"{row.message}"
                ),
            )

        if len(rejected_rows) > 10:
            messages.error(
                request,
                (
                    f"يوجد {len(rejected_rows) - 10} "
                    "صفوف مرفوضة إضافية لم تُعرض."
                ),
            )

    for warning in result.warnings[:10]:
        messages.warning(
            request,
            warning,
        )

    if len(result.warnings) > 10:
        messages.warning(
            request,
            (
                f"يوجد {len(result.warnings) - 10} "
                "تنبيهات إضافية لم تُعرض."
            ),
        )

    messages.info(
        request,
        (
            "ملخص الاستيراد: "
            f"أضيف {result.added}، "
            f"حُدّث {result.updated}، "
            f"دون تغيير {result.unchanged}، "
            f"مرفوض {result.rejected}."
        ),
    )

    return redirect(
        "municipal_portal:services_assets",
    )
@municipal_required
def asset_excel_template(request):
    workbook = Workbook()
    worksheet = workbook.active

    worksheet.title = "الأصول البلدية"
    worksheet.sheet_view.rightToLeft = True
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:L200"

    headers = [
        "رمز الأصل *",
        "الخدمة *",
        "اسم الأصل",
        "الحي",
        "العنوان",
        "حالة التشغيل *",
        "الحالة الفنية *",
        "تاريخ التركيب",
        "تاريخ آخر صيانة",
        "خط العرض *",
        "خط الطول *",
        "ملاحظات",
    ]

    worksheet.append(headers)

    # صف توضيحي يمكن للمستخدم حذفه
    worksheet.append(
        [
            "LIGHT-001",
            "الإنارة العامة",
            "إنارة حي النهضة",
            "",
            "حي النهضة - شارع الزبيري",
            "يعمل",
            "جيدة",
            "2026-01-15",
            "2026-09-20",
            15.3694000,
            44.1910000,
            "",
        ]
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="218653",
    )

    light_green_fill = PatternFill(
        fill_type="solid",
        fgColor="EAF5EF",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
        size=12,
    )

    thin_green = Side(
        style="thin",
        color="CFE3D7",
    )

    border = Border(
        left=thin_green,
        right=thin_green,
        top=thin_green,
        bottom=thin_green,
    )

    for cell in worksheet[1]:
        cell.fill = green_fill
        cell.font = white_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = border

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=200,
        min_col=1,
        max_col=len(headers),
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
            )
            cell.border = border

    for cell in worksheet[2]:
        cell.fill = light_green_fill

    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 26

    widths = {
        "A": 18,
        "B": 24,
        "C": 28,
        "D": 22,
        "E": 35,
        "F": 22,
        "G": 20,
        "H": 20,
        "I": 22,
        "J": 18,
        "K": 18,
        "L": 35,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    # قائمة الخدمات الموجودة فعليًا في قاعدة البيانات
    service_names = list(
        MunicipalService.objects
        .order_by("name_ar")
        .values_list("name_ar", flat=True)
    )

    if service_names:
        service_formula = '"{}"'.format(
            ",".join(service_names)
        )

        service_validation = DataValidation(
            type="list",
            formula1=service_formula,
            allow_blank=False,
        )

        worksheet.add_data_validation(
            service_validation
        )

        service_validation.add("B2:B200")

    status_validation = DataValidation(
        type="list",
        formula1='"يعمل,يحتاج صيانة,تحت الصيانة,خارج الخدمة"',
        allow_blank=False,
    )

    condition_validation = DataValidation(
        type="list",
        formula1='"جيدة,مقبولة,سيئة,حرجة"',
        allow_blank=False,
    )

    worksheet.add_data_validation(
        status_validation
    )

    worksheet.add_data_validation(
        condition_validation
    )

    status_validation.add("F2:F200")
    condition_validation.add("G2:G200")

    for row_number in range(2, 201):
        worksheet[
            f"H{row_number}"
        ].number_format = "yyyy-mm-dd"

        worksheet[
            f"I{row_number}"
        ].number_format = "yyyy-mm-dd"

        worksheet[
            f"J{row_number}"
        ].number_format = "0.0000000"

        worksheet[
            f"K{row_number}"
        ].number_format = "0.0000000"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        'attachment; '
        'filename="cityvoice_assets_ar.xlsx"'
    )

    return response

@municipal_required
def complaint_list(request):
    complaints = Complaint.objects.select_related(
        "citizen",
        "service",
        "category",
        "neighborhood",
    )

    search_query = request.GET.get(
        "q",
        "",
    ).strip()

    selected_status = request.GET.get(
        "status",
        "",
    ).strip()

    selected_priority = request.GET.get(
        "priority",
        "",
    ).strip()

    selected_service = request.GET.get(
        "service",
        "",
    ).strip()

    if search_query:
        complaints = complaints.filter(
            Q(
                reference_number__icontains=search_query,
            )
            | Q(
                description__icontains=search_query,
            )
            | Q(
                address__icontains=search_query,
            )
            | Q(
                citizen__username__icontains=search_query,
            )
            | Q(
                citizen__first_name__icontains=search_query,
            )
            | Q(
                citizen__last_name__icontains=search_query,
            )
            | Q(
                service__name_ar__icontains=search_query,
            )
            | Q(
                category__name_ar__icontains=search_query,
            )
        )

    if selected_status:
        complaints = complaints.filter(
            status=selected_status,
        )

    if selected_priority:
        complaints = complaints.filter(
            priority=selected_priority,
        )

    if selected_service:
        complaints = complaints.filter(
            service_id=selected_service,
        )

    context = {
        "complaints": complaints.order_by("-created_at"),

        "services": (
            MunicipalService.objects
            .all()
            .order_by("name_ar")
        ),

        "status_choices": Complaint.Status.choices,
        "priority_choices": Complaint.Priority.choices,

        "search_query": search_query,
        "selected_status": selected_status,
        "selected_priority": selected_priority,
        "selected_service": selected_service,
    }

    return render(
        request,
        "municipal_portal/complaints/list.html",
        context,
    )


@municipal_required
def complaint_detail(request, reference_number):
    complaint = get_object_or_404(
        Complaint.objects.select_related(
            "citizen",
            "service",
            "category",
            "asset",
            "neighborhood",
        ),
        reference_number=reference_number,
    )

    try:
        work_order = complaint.work_order
    except WorkOrder.DoesNotExist:
        work_order = None

    context = {
        "complaint": complaint,
        "work_order": work_order,
        "maptiler_api_key": settings.MAPTILER_API_KEY,
    }

    return render(
        request,
        "municipal_portal/complaints/detail.html",
        context,
    )


@municipal_required
def work_order_list(request):
    work_orders = WorkOrder.objects.select_related(
        "complaint",
        "complaint__service",
        "assigned_to",
        "assigned_by",
    )

    search_query = request.GET.get(
        "q",
        "",
    ).strip()

    selected_status = request.GET.get(
        "status",
        "",
    ).strip()

    if search_query:
        work_orders = work_orders.filter(
            Q(
                complaint__reference_number__icontains=
                search_query
            )
            | Q(
                assigned_to__username__icontains=
                search_query
            )
            | Q(
                assigned_to__first_name__icontains=
                search_query
            )
            | Q(
                assigned_to__last_name__icontains=
                search_query
            )
            | Q(
                instructions__icontains=search_query
            )
        )

    if selected_status:
        work_orders = work_orders.filter(
            status=selected_status,
        )

    context = {
        "work_orders": work_orders.order_by(
            "-created_at",
        ),

        "status_choices": WorkOrder.Status.choices,
        "search_query": search_query,
        "selected_status": selected_status,
    }

    return render(
        request,
        "municipal_portal/work_orders/list.html",
        context,
    )


@municipal_required
def worker_list(request):
    workers = (
        User.objects
        .filter(
            role=User.Role.FIELD_WORKER,
            is_active=True,
        )
        .annotate(
            total_orders=Count(
                "assigned_work_orders",
                distinct=True,
            )
        )
        .order_by(
            "first_name",
            "last_name",
            "username",
        )
    )

    return render(
        request,
        "municipal_portal/workers/list.html",
        {
            "workers": workers,
        },
    )


@municipal_required
def notification_list(request):
    notifications = (
        Notification.objects
        .filter(
            recipient=request.user,
        )
        .order_by("-created_at")
    )

    return render(
        request,
        "municipal_portal/notifications.html",
        {
            "notifications": notifications,
        },
    )


@municipal_required
@require_POST
def mark_notification_read(request, notification_id):
    notification = get_object_or_404(
        Notification,
        pk=notification_id,
        recipient=request.user,
    )

    if not notification.is_read:
        notification.is_read = True
        notification.read_at = timezone.now()

        notification.save(
            update_fields=[
                "is_read",
                "read_at",
            ]
        )

    if notification.link:
        return redirect(notification.link)

    return redirect(
        "municipal_portal:notification_list",
    )


@municipal_required
@require_POST
def mark_all_notifications_read(request):
    Notification.objects.filter(
        recipient=request.user,
        is_read=False,
    ).update(
        is_read=True,
        read_at=timezone.now(),
    )

    messages.success(
        request,
        "تم تحديد جميع الإشعارات كمقروءة.",
    )

    return redirect(
        "municipal_portal:notification_list",
    )