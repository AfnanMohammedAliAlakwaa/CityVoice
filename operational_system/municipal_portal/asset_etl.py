import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.contrib.gis.geos import Point
from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from assets.models import AssetType, MunicipalAsset
from locations.models import Neighborhood
from municipal_services.models import MunicipalService


# =========================================================
# نتيجة معالجة صف واحد
# =========================================================

@dataclass
class RowResult:
    row_number: int
    status: str
    asset_code: str = ""
    message: str = ""


# =========================================================
# نتيجة ملف الاستيراد كاملًا
# =========================================================

@dataclass
class ETLResult:
    added: int = 0
    updated: int = 0
    unchanged: int = 0
    rejected: int = 0
    empty_rows: int = 0
    warnings: list[str] = field(default_factory=list)
    rows: list[RowResult] = field(default_factory=list)

    @property
    def processed(self):
        return (
            self.added
            + self.updated
            + self.unchanged
            + self.rejected
        )

    @property
    def successful(self):
        return self.added + self.updated + self.unchanged

    def add_row(
        self,
        row_number,
        status,
        asset_code="",
        message="",
    ):
        self.rows.append(
            RowResult(
                row_number=row_number,
                status=status,
                asset_code=asset_code,
                message=message,
            )
        )

        if status == "added":
            self.added += 1
        elif status == "updated":
            self.updated += 1
        elif status == "unchanged":
            self.unchanged += 1
        elif status == "rejected":
            self.rejected += 1


# =========================================================
# خطأ متعلق بالملف كاملًا
# =========================================================

class AssetETLError(Exception):
    pass


# =========================================================
# أسماء الحقول والمرادفات المقبولة
# =========================================================

HEADER_ALIASES = {
    "asset_code": {
        "رمز الأصل",
        "رمز الاصل",
        "كود الأصل",
        "كود الاصل",
        "رقم الأصل",
        "رقم الاصل",
        "asset code",
        "asset_code",
        "asset id",
        "asset_id",
        "code",
        "id",
    },

    "service": {
        "الخدمة",
        "اسم الخدمة",
        "نوع الخدمة",
        "service",
        "service name",
        "service_name",
        "category",
    },

    "name_ar": {
        "اسم الأصل",
        "اسم الاصل",
        "الاسم",
        "وصف الأصل",
        "وصف الاصل",
        "asset name",
        "asset_name",
        "name",
        "name ar",
        "name_ar",
    },

    "neighborhood": {
        "الحي",
        "اسم الحي",
        "المنطقة",
        "neighborhood",
        "neighbourhood",
        "district",
        "area",
    },

    "address": {
        "العنوان",
        "الموقع",
        "وصف الموقع",
        "عنوان الأصل",
        "عنوان الاصل",
        "address",
        "location description",
        "site",
    },

    "status": {
        "حالة التشغيل",
        "الحالة التشغيلية",
        "التشغيل",
        "status",
        "operational status",
        "operating status",
        "asset status",
    },

    "condition": {
        "الحالة الفنية",
        "حالة الأصل",
        "حالة الاصل",
        "التقييم الفني",
        "condition",
        "technical condition",
        "asset condition",
    },

    "installation_date": {
        "تاريخ التركيب",
        "تاريخ الإنشاء",
        "تاريخ الانشاء",
        "تاريخ التشغيل",
        "installation date",
        "installation_date",
        "installed at",
    },

    "last_maintenance_date": {
        "تاريخ آخر صيانة",
        "تاريخ اخر صيانة",
        "آخر صيانة",
        "اخر صيانة",
        "last maintenance",
        "last maintenance date",
        "last_maintenance_date",
        "maintenance date",
    },

    "latitude": {
        "خط العرض",
        "احداثي العرض",
        "إحداثي العرض",
        "الإحداثي y",
        "الاحداثي y",
        "latitude",
        "lat",
        "y",
    },

    "longitude": {
        "خط الطول",
        "احداثي الطول",
        "إحداثي الطول",
        "الإحداثي x",
        "الاحداثي x",
        "longitude",
        "long",
        "lng",
        "lon",
        "x",
    },

    "notes": {
        "ملاحظات",
        "الملاحظات",
        "ملاحظة",
        "تفاصيل",
        "notes",
        "note",
        "comments",
        "comment",
    },
}


# الحقول التي لا يمكن إنشاء أصل صالح بدونها.
REQUIRED_FIELDS = {
    "asset_code",
    "service",
    "latitude",
    "longitude",
}


# =========================================================
# حالات التشغيل والحالة الفنية
# =========================================================

STATUS_ALIASES = {
    "يعمل": MunicipalAsset.Status.OPERATIONAL,
    "فعال": MunicipalAsset.Status.OPERATIONAL,
    "نشط": MunicipalAsset.Status.OPERATIONAL,
    "سليم": MunicipalAsset.Status.OPERATIONAL,
    "operational": MunicipalAsset.Status.OPERATIONAL,
    "active": MunicipalAsset.Status.OPERATIONAL,
    "working": MunicipalAsset.Status.OPERATIONAL,

    "يحتاج صيانة": MunicipalAsset.Status.NEEDS_MAINTENANCE,
    "بحاجة إلى صيانة": MunicipalAsset.Status.NEEDS_MAINTENANCE,
    "بحاجة لصيانة": MunicipalAsset.Status.NEEDS_MAINTENANCE,
    "needs maintenance": MunicipalAsset.Status.NEEDS_MAINTENANCE,
    "maintenance required": MunicipalAsset.Status.NEEDS_MAINTENANCE,

    "تحت الصيانة": MunicipalAsset.Status.UNDER_MAINTENANCE,
    "قيد الصيانة": MunicipalAsset.Status.UNDER_MAINTENANCE,
    "under maintenance": MunicipalAsset.Status.UNDER_MAINTENANCE,

    "خارج الخدمة": MunicipalAsset.Status.OUT_OF_SERVICE,
    "متوقف": MunicipalAsset.Status.OUT_OF_SERVICE,
    "لا يعمل": MunicipalAsset.Status.OUT_OF_SERVICE,
    "out of service": MunicipalAsset.Status.OUT_OF_SERVICE,
    "inactive": MunicipalAsset.Status.OUT_OF_SERVICE,
    "not working": MunicipalAsset.Status.OUT_OF_SERVICE,
    "غير محدد": MunicipalAsset.Status.UNKNOWN,
"غير معروف": MunicipalAsset.Status.UNKNOWN,
"not defined": MunicipalAsset.Status.UNKNOWN,
"unknown": MunicipalAsset.Status.UNKNOWN,
"n/a": MunicipalAsset.Status.UNKNOWN,
}


CONDITION_ALIASES = {
    "جيدة": MunicipalAsset.Condition.GOOD,
    "جيد": MunicipalAsset.Condition.GOOD,
    "سليمة": MunicipalAsset.Condition.GOOD,
    "سليم": MunicipalAsset.Condition.GOOD,
    "good": MunicipalAsset.Condition.GOOD,

    "مقبولة": MunicipalAsset.Condition.FAIR,
    "مقبول": MunicipalAsset.Condition.FAIR,
    "متوسطة": MunicipalAsset.Condition.FAIR,
    "متوسط": MunicipalAsset.Condition.FAIR,
    "fair": MunicipalAsset.Condition.FAIR,
    "average": MunicipalAsset.Condition.FAIR,

    "سيئة": MunicipalAsset.Condition.POOR,
    "سيء": MunicipalAsset.Condition.POOR,
    "ضعيفة": MunicipalAsset.Condition.POOR,
    "ضعيف": MunicipalAsset.Condition.POOR,
    "poor": MunicipalAsset.Condition.POOR,
    "bad": MunicipalAsset.Condition.POOR,

    "حرجة": MunicipalAsset.Condition.CRITICAL,
    "حرج": MunicipalAsset.Condition.CRITICAL,
    "خطرة": MunicipalAsset.Condition.CRITICAL,
    "critical": MunicipalAsset.Condition.CRITICAL,
    "غير محددة": MunicipalAsset.Condition.UNKNOWN,
"غير محدد": MunicipalAsset.Condition.UNKNOWN,
"غير معروفة": MunicipalAsset.Condition.UNKNOWN,
"غير معروف": MunicipalAsset.Condition.UNKNOWN,
"not defined": MunicipalAsset.Condition.UNKNOWN,
"unknown": MunicipalAsset.Condition.UNKNOWN,
"n/a": MunicipalAsset.Condition.UNKNOWN,
}


ARABIC_DIGITS = str.maketrans(
    "٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹",
    "01234567890123456789",
)


# =========================================================
# أدوات التنظيف والتحويل
# =========================================================

def clean_text(value):
    if value is None:
        return ""

    text = str(value)
    text = text.translate(ARABIC_DIGITS)
    text = text.replace("\u200f", "")
    text = text.replace("\u200e", "")
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def normalize_text(value):
    text = clean_text(value).casefold()

    text = text.replace("أ", "ا")
    text = text.replace("إ", "ا")
    text = text.replace("آ", "ا")
    text = text.replace("ى", "ي")
    text = text.replace("ة", "ه")

    text = re.sub(r"[_\-]+", " ", text)
    text = re.sub(r"[^\w\s\u0600-\u06ff]", "", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def normalize_header(value):
    text = clean_text(value)

    # حذف النجمة المستخدمة للدلالة على الحقل الإلزامي.
    text = text.replace("*", "")
    text = text.replace("：", ":")
    text = text.rstrip(":")

    return normalize_text(text)


def normalize_asset_code(value):
    value = clean_text(value).upper()
    value = re.sub(r"\s+", "-", value)
    return value


def parse_decimal(value):
    if value in (None, ""):
        return None

    if isinstance(value, (int, float, Decimal)):
        return float(value)

    text = clean_text(value)
    text = text.replace("٬", "")
    text = text.replace(",", ".")
    text = text.replace("٫", ".")

    # يسمح بقيمة مثل: 15.3694° N
    match = re.search(
        r"[-+]?\d+(?:\.\d+)?",
        text,
    )

    if not match:
        raise ValueError("القيمة ليست رقمًا.")

    try:
        return float(
            Decimal(match.group())
        )
    except (InvalidOperation, ValueError):
        raise ValueError("القيمة ليست رقمًا.")


def parse_date_value(value, workbook_epoch=None):
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    # بعض ملفات Excel تخزن التاريخ كرقم تسلسلي.
    if isinstance(value, (int, float)):
        if workbook_epoch is None:
            raise ValueError("تعذر تفسير التاريخ الرقمي.")

        return from_excel(
            value,
            workbook_epoch,
        ).date()

    text = clean_text(value)

    date_formats = [
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d.%m.%Y",
        "%Y.%m.%d",
    ]

    for date_format in date_formats:
        try:
            return datetime.strptime(
                text,
                date_format,
            ).date()
        except ValueError:
            continue

    raise ValueError(
        "صيغة التاريخ غير معروفة."
    )


def normalized_aliases(alias_mapping):
    result = {}

    for field_name, aliases in alias_mapping.items():
        for alias in aliases:
            result[normalize_header(alias)] = field_name

    return result


NORMALIZED_HEADER_ALIASES = normalized_aliases(
    HEADER_ALIASES
)

NORMALIZED_STATUS_ALIASES = {
    normalize_text(key): value
    for key, value in STATUS_ALIASES.items()
}

NORMALIZED_CONDITION_ALIASES = {
    normalize_text(key): value
    for key, value in CONDITION_ALIASES.items()
}


# =========================================================
# مطابقة الأعمدة
# =========================================================

def build_column_map(worksheet):
    column_map = {}
    unknown_headers = []
    duplicate_fields = []

    for column_number, cell in enumerate(
        worksheet[1],
        start=1,
    ):
        original_header = clean_text(cell.value)

        if not original_header:
            continue

        normalized_header = normalize_header(
            original_header
        )

        field_name = NORMALIZED_HEADER_ALIASES.get(
            normalized_header
        )

        if field_name is None:
            unknown_headers.append(
                original_header
            )
            continue

        if field_name in column_map:
            duplicate_fields.append(
                original_header
            )
            continue

        column_map[field_name] = column_number

    missing_fields = (
        REQUIRED_FIELDS - set(column_map)
    )

    if missing_fields:
        arabic_names = {
            "asset_code": "رمز الأصل",
            "service": "الخدمة",
            "latitude": "خط العرض",
            "longitude": "خط الطول",
        }

        missing_names = [
            arabic_names[field_name]
            for field_name in sorted(missing_fields)
        ]

        raise AssetETLError(
            "الملف لا يحتوي على الأعمدة الأساسية التالية: "
            + "، ".join(missing_names)
            + "."
        )

    if duplicate_fields:
        raise AssetETLError(
            "توجد أعمدة متكررة أو تشير إلى الحقل نفسه: "
            + "، ".join(duplicate_fields)
            + "."
        )

    return column_map, unknown_headers


def get_row_value(row, column_map, field_name):
    column_number = column_map.get(field_name)

    if column_number is None:
        return None

    index = column_number - 1

    if index >= len(row):
        return None

    return row[index]


# =========================================================
# البحث في البيانات المرجعية
# =========================================================

def find_service(service_name):
    service_name = clean_text(service_name)

    if not service_name:
        return None

    service = (
        MunicipalService.objects
        .filter(is_active=True)
        .filter(
            Q(name_ar__iexact=service_name)
            | Q(name_en__iexact=service_name)
        )
        .first()
    )

    if service:
        return service

    wanted = normalize_text(service_name)

    for candidate in MunicipalService.objects.filter(
        is_active=True
    ):
        if wanted in {
            normalize_text(candidate.name_ar),
            normalize_text(candidate.name_en),
        }:
            return candidate

    return None


def find_neighborhood(neighborhood_name):
    neighborhood_name = clean_text(
        neighborhood_name
    )

    if not neighborhood_name:
        return None

    neighborhood = (
        Neighborhood.objects
        .filter(is_active=True)
        .filter(
            Q(name_ar__iexact=neighborhood_name)
            | Q(name_en__iexact=neighborhood_name)
        )
        .first()
    )

    if neighborhood:
        return neighborhood

    wanted = normalize_text(neighborhood_name)

    for candidate in Neighborhood.objects.filter(
        is_active=True
    ):
        if wanted in {
            normalize_text(candidate.name_ar),
            normalize_text(candidate.name_en),
        }:
            return candidate

    return None


def find_default_asset_type(service):
    return (
        AssetType.objects
        .filter(
            service=service,
            is_active=True,
        )
        .order_by("id")
        .first()
    )


# =========================================================
# معالجة الإحداثيات
# =========================================================

def clean_coordinates(latitude, longitude):
    latitude = parse_decimal(latitude)
    longitude = parse_decimal(longitude)

    if latitude is None or longitude is None:
        raise ValueError(
            "خط العرض وخط الطول مطلوبان."
        )

    swapped = False

    # إذا كانت قيمة العرض خارج نطاقه، ولكنها صالحة كطول،
    # وكانت قيمة الطول صالحة كعرض، يتم تبديلهما.
    if (
        not -90 <= latitude <= 90
        and -180 <= latitude <= 180
        and -90 <= longitude <= 90
    ):
        latitude, longitude = longitude, latitude
        swapped = True

    if not -90 <= latitude <= 90:
        raise ValueError(
            "خط العرض يجب أن يكون بين -90 و90."
        )

    if not -180 <= longitude <= 180:
        raise ValueError(
            "خط الطول يجب أن يكون بين -180 و180."
        )

    return latitude, longitude, swapped


# =========================================================
# مقارنة بيانات الأصل
# =========================================================

def asset_has_changes(asset, values):
    current_location = asset.location

    current_latitude = (
        round(current_location.y, 7)
        if current_location is not None
        else None
    )

    current_longitude = (
        round(current_location.x, 7)
        if current_location is not None
        else None
    )

    new_neighborhood_id = (
        values["neighborhood"].pk
        if values["neighborhood"]
        else None
    )

    comparisons = [
        asset.asset_type_id
        != values["asset_type"].pk,

        clean_text(asset.name_ar)
        != values["name_ar"],

        asset.neighborhood_id
        != new_neighborhood_id,

        clean_text(asset.address)
        != values["address"],

        asset.status
        != values["status"],

        asset.condition
        != values["condition"],

        asset.installation_date
        != values["installation_date"],

        asset.last_maintenance_date
        != values["last_maintenance_date"],

        clean_text(asset.notes)
        != values["notes"],

        current_latitude
        != round(values["latitude"], 7),

        current_longitude
        != round(values["longitude"], 7),
    ]

    return any(comparisons)


# =========================================================
# إنشاء الأصل أو تحديثه
# =========================================================

def save_asset(asset_code, values):
    asset = (
        MunicipalAsset.objects
        .filter(
            asset_code__iexact=asset_code
        )
        .first()
    )

    if asset is None:
        MunicipalAsset.objects.create(
            asset_code=asset_code,
            asset_type=values["asset_type"],
            name_ar=values["name_ar"],
            neighborhood=values["neighborhood"],
            address=values["address"],
            location=Point(
                values["longitude"],
                values["latitude"],
                srid=4326,
            ),
            status=values["status"],
            condition=values["condition"],
            installation_date=(
                values["installation_date"]
            ),
            last_maintenance_date=(
                values["last_maintenance_date"]
            ),
            notes=values["notes"],
        )

        return "added"

    if not asset_has_changes(asset, values):
        return "unchanged"

    asset.asset_code = asset_code
    asset.asset_type = values["asset_type"]
    asset.name_ar = values["name_ar"]
    asset.neighborhood = values["neighborhood"]
    asset.address = values["address"]
    asset.location = Point(
        values["longitude"],
        values["latitude"],
        srid=4326,
    )
    asset.status = values["status"]
    asset.condition = values["condition"]
    asset.installation_date = (
        values["installation_date"]
    )
    asset.last_maintenance_date = (
        values["last_maintenance_date"]
    )
    asset.notes = values["notes"]

    asset.save()

    return "updated"


# =========================================================
# محرك ETL الرئيسي
# =========================================================

def process_asset_excel(uploaded_file):
    result = ETLResult()

    try:
        workbook = load_workbook(
            uploaded_file,
            data_only=True,
            read_only=True,
        )
    except Exception as error:
        raise AssetETLError(
            "تعذر قراءة ملف Excel. تأكدي من أن الملف سليم وبصيغة xlsx."
        ) from error

    worksheet = workbook.active

    first_row = next(
    worksheet.iter_rows(
        min_row=1,
        max_row=1,
        values_only=True,
    ),
    None,
)

    if first_row is None or not any(
        value not in (None, "")
        for value in first_row
):
        raise AssetETLError(
            "ملف Excel فارغ أو لا يحتوي على أسماء أعمدة."
    )

    column_map, unknown_headers = build_column_map(
        worksheet
    )

    if unknown_headers:
        result.warnings.append(
            "تم تجاهل الأعمدة غير المعروفة: "
            + "، ".join(unknown_headers)
            + "."
        )

    codes_seen = set()

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        if not any(
            value not in (None, "")
            for value in row
        ):
            result.empty_rows += 1
            continue

        asset_code = normalize_asset_code(
            get_row_value(
                row,
                column_map,
                "asset_code",
            )
        )

        try:
            if not asset_code:
                raise ValueError(
                    "رمز الأصل مطلوب."
                )

            normalized_code = asset_code.casefold()

            if normalized_code in codes_seen:
                raise ValueError(
                    (
                        f"رمز الأصل «{asset_code}» "
                        "مكرر داخل ملف Excel."
                    )
                )

            codes_seen.add(normalized_code)

            service_name = clean_text(
                get_row_value(
                    row,
                    column_map,
                    "service",
                )
            )

            if not service_name:
                raise ValueError(
                    "الخدمة مطلوبة."
                )

            service = find_service(
                service_name
            )

            if service is None:
                raise ValueError(
                    (
                        f"الخدمة «{service_name}» "
                        "غير موجودة أو غير مفعلة."
                    )
                )

            asset_type = find_default_asset_type(
                service
            )

            if asset_type is None:
                raise ValueError(
                    (
                        "لا يوجد نوع أصل داخلي مفعّل "
                        f"مرتبط بخدمة «{service.name_ar}»."
                    )
                )

            latitude, longitude, swapped = (
                clean_coordinates(
                    get_row_value(
                        row,
                        column_map,
                        "latitude",
                    ),
                    get_row_value(
                        row,
                        column_map,
                        "longitude",
                    ),
                )
            )

            if swapped:
                result.warnings.append(
                    (
                        f"الصف {row_number}: تم تبديل خط العرض "
                        "وخط الطول لأنهما كانا في موضعين معكوسين."
                    )
                )

            neighborhood_name = clean_text(
                get_row_value(
                    row,
                    column_map,
                    "neighborhood",
                )
            )

            neighborhood = None

            if neighborhood_name:
                neighborhood = find_neighborhood(
                    neighborhood_name
                )

                if neighborhood is None:
                    result.warnings.append(
                        (
                            f"الصف {row_number}: الحي "
                            f"«{neighborhood_name}» غير موجود؛ "
                            "تم استيراد الأصل دون ربطه بحي."
                        )
                    )

            status_value = clean_text(
                get_row_value(
                    row,
                    column_map,
                    "status",
                )
            )

            if status_value:
                status = NORMALIZED_STATUS_ALIASES.get(
                    normalize_text(status_value)
                )

                if status is None:
                    raise ValueError(
                        (
                            f"حالة التشغيل «{status_value}» "
                            "غير معروفة."
                        )
                    )
            else:
                status = MunicipalAsset.Status.UNKNOWN

                result.warnings.append(
                    (
                        f"الصف {row_number}: حالة التشغيل فارغة؛ "
                        "تم تسجيلها «غير محدد»."
                    )
                )

            condition_value = clean_text(
                get_row_value(
                    row,
                    column_map,
                    "condition",
                )
            )

            if condition_value:
                condition = (
                    NORMALIZED_CONDITION_ALIASES.get(
                        normalize_text(
                            condition_value
                        )
                    )
                )

                if condition is None:
                    raise ValueError(
                        (
                            f"الحالة الفنية «{condition_value}» "
                            "غير معروفة."
                        )
                    )
            else:
                condition = MunicipalAsset.Condition.UNKNOWN

                result.warnings.append(
                    (
                        f"الصف {row_number}: الحالة الفنية فارغة؛ "
                        "تم تسجيلها «غير محددة»."
                    )
                )

            installation_date = parse_date_value(
                get_row_value(
                    row,
                    column_map,
                    "installation_date",
                ),
                workbook.epoch,
            )

            last_maintenance_date = parse_date_value(
                get_row_value(
                    row,
                    column_map,
                    "last_maintenance_date",
                ),
                workbook.epoch,
            )

            values = {
                "asset_type": asset_type,

                "name_ar": clean_text(
                    get_row_value(
                        row,
                        column_map,
                        "name_ar",
                    )
                ),

                "neighborhood": neighborhood,

                "address": clean_text(
                    get_row_value(
                        row,
                        column_map,
                        "address",
                    )
                ),

                "status": status,
                "condition": condition,

                "installation_date": (
                    installation_date
                ),

                "last_maintenance_date": (
                    last_maintenance_date
                ),

                "latitude": latitude,
                "longitude": longitude,

                "notes": clean_text(
                    get_row_value(
                        row,
                        column_map,
                        "notes",
                    )
                ),
            }

            with transaction.atomic():
                operation = save_asset(
                    asset_code,
                    values,
                )

            operation_messages = {
                "added": "تمت إضافة الأصل.",
                "updated": "تم تحديث الأصل.",
                "unchanged": (
                    "لم تتغير بيانات الأصل."
                ),
            }

            result.add_row(
                row_number=row_number,
                status=operation,
                asset_code=asset_code,
                message=operation_messages[operation],
            )

        except (
            ValueError,
            TypeError,
            InvalidOperation,
        ) as error:
            result.add_row(
                row_number=row_number,
                status="rejected",
                asset_code=asset_code,
                message=str(error),
            )

        except Exception as error:
            result.add_row(
                row_number=row_number,
                status="rejected",
                asset_code=asset_code,
                message=(
                    "حدث خطأ غير متوقع أثناء معالجة الصف: "
                    f"{error}"
                ),
            )

    workbook.close()

    return result